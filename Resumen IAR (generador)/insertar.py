# -*- coding: utf-8 -*-
"""Inserta contenido nuevo en un .docx YA EXISTENTE sin reconstruirlo.

El usuario editó el documento a mano, asi que no se puede regenerar desde
build.py. python-docx solo sabe agregar al final; para insertar en el medio
se crea el contenido al final del documento DESTINO (asi las imagenes quedan
con su relacion r:embed valida en ese archivo) y despues se mueven los
elementos XML recien creados al lugar que corresponde.
"""
import re
import docx


class Doc:
    def __init__(self, ruta):
        self.ruta = ruta
        self.doc = docx.Document(ruta)
        self.body = self.doc.element.body

    # ---------- localizar puntos de anclaje ----------
    def headings(self):
        """[(indice_parrafo, nivel, titulo, parrafo)] de todos los encabezados."""
        out = []
        for i, p in enumerate(self.doc.paragraphs):
            m = re.match(r"Heading (\d)", p.style.name or "")
            if m and p.text.strip():
                out.append((i, int(m.group(1)), p.text.strip(), p))
        return out

    def heading(self, texto):
        """Encabezado cuyo titulo empieza con `texto`. Error si no es unico."""
        c = [h for h in self.headings() if h[2].startswith(texto)]
        if len(c) != 1:
            raise KeyError(f"{texto!r} coincide con {len(c)} encabezados: "
                           f"{[x[2] for x in c]}")
        return c[0]

    def parrafo(self, texto):
        """Parrafo cualquiera (no encabezado) que empieza con `texto`. Sirve
        para anclar en el medio de una seccion, p. ej. antes de un ejemplo."""
        c = [p for p in self.doc.paragraphs if p.text.strip().startswith(texto)]
        if len(c) != 1:
            raise KeyError(f"{texto!r} coincide con {len(c)} parrafos")
        return c[0]

    def fin_de_seccion(self, texto):
        """Parrafo ANTE EL CUAL insertar para quedar al final de esa seccion:
        el siguiente encabezado de nivel <= al de la seccion."""
        i, lvl, _t, _p = self.heading(texto)
        for j, l2, _t2, p2 in self.headings():
            if j > i and l2 <= lvl:
                return p2
        return None      # la seccion llega hasta el final del documento

    # ---------- insercion ----------
    def insertar_antes(self, ancla, escribir):
        """Ejecuta `escribir(doc)` —que agrega al final— y mueve lo agregado
        justo antes del parrafo `ancla`.

        Se identifica lo nuevo por posicion, NO por id(): los proxies de lxml
        se crean bajo demanda y su id se reutiliza, asi que un set de id()
        da falsos positivos. python-docx inserta siempre antes del <w:sectPr>
        final, de modo que lo agregado queda en body[n_antes-1:-1]."""
        assert self.body[-1].tag.endswith("}sectPr"), "el body no termina en sectPr"
        n_antes = len(self.body)
        escribir(self.doc)
        nuevos = list(self.body[n_antes - 1:-1])
        if ancla is None:
            return len(nuevos)          # ya quedo al final, no hay que mover
        for e in nuevos:
            ancla._p.addprevious(e)
        return len(nuevos)

    # ---------- mantenimiento ----------
    def renumerar_figuras(self):
        """Las leyendas son texto plano, asi que al insertar una figura en el
        medio hay que renumerar todas las de abajo."""
        n = 0
        for p in self.doc.paragraphs:
            if re.match(r"^Figura \d+ — ", p.text):
                n += 1
                for r in p.runs:
                    m = re.match(r"^(Figura )(\d+)( — .*)$", r.text, re.S)
                    if m:
                        r.text = f"{m.group(1)}{n}{m.group(3)}"
                        break
                else:
                    # la leyenda quedo partida en varios runs
                    txt = re.sub(r"^Figura \d+", f"Figura {n}", p.text)
                    for r in p.runs[1:]:
                        r.text = ""
                    p.runs[0].text = txt
        return n

    def guardar(self, ruta=None):
        self.doc.save(ruta or self.ruta)
